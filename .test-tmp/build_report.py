from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r'C:\Users\gabri\OneDrive\Desktop\ideia-tech\Ideia-Tech\Relatorio_Auditoria_MVP_IdeiaTech_2026-10-22.docx'
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.62)
sec.left_margin = Inches(.72); sec.right_margin = Inches(.72)

navy = '17324D'; blue = '167D9A'; pale = 'EAF2F6'; gray = '536273'; red = 'A33A36'; amber = 'A36A12'; green = '2E6E52'
styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(9.3); styles['Normal'].font.color.rgb = RGBColor.from_string('253443')
styles['Normal'].paragraph_format.space_after = Pt(4)
for name,size,color in [('Title',26,'000000'),('Heading 1',16,navy),('Heading 2',11,blue)]:
    s=styles[name]; s.font.name='Aptos Display'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color)
    s.paragraph_format.space_before=Pt(9 if name!='Title' else 0); s.paragraph_format.space_after=Pt(4)

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def set_cell(cell, text, bold=False, color=None, size=8.2):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(1); p.paragraph_format.space_before=Pt(1)
    r=p.add_run(text); r.bold=bold; r.font.name='Aptos'; r.font.size=Pt(size)
    if color: r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'; t.autofit=False
    for i,h in enumerate(headers): set_cell(t.rows[0].cells[i],h,True,'FFFFFF',8); shade(t.rows[0].cells[i],navy)
    trPr=t.rows[0]._tr.get_or_add_trPr(); rep=OxmlElement('w:tblHeader'); rep.set(qn('w:val'),'true'); trPr.append(rep)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            set_cell(cells[i],val)
            if ri%2==1: shade(cells[i],'F3F6F8')
    if widths:
        for row in t.rows:
            for cell,w in zip(row.cells,widths): cell.width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t
def para(text='', style=None, boldlead=None):
    p=doc.add_paragraph(style=style)
    if boldlead and text.startswith(boldlead):
        p.add_run(boldlead).bold=True; p.add_run(text[len(boldlead):])
    else: p.add_run(text)
    return p
def bullet(text):
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(2); p.add_run(text)
    return p

p=doc.add_paragraph(style='Title');
# Explicitly suppress the built-in Title style's inherited bottom rule.
pPr=p._p.get_or_add_pPr(); borders=OxmlElement('w:pBdr'); bottom=OxmlElement('w:bottom'); bottom.set(qn('w:val'),'nil'); borders.append(bottom); pPr.append(borders)
p.add_run('Auditoria do sistema IdeiaTech')
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(9)
r=p.add_run('Relatório de prontidão para apresentação do MVP'); r.font.size=Pt(13); r.font.color.rgb=RGBColor.from_string(blue); r.bold=True
para('Avaliação funcional, técnica, estrutural e visual para apoiar a preparação da demonstração prevista para 23 de outubro de 2026 à tarde.')
table(['RESULTADO','EVIDÊNCIA','DECISÃO PARA A DEMO'], [[
    'Adequado para demonstração local guiada, com ressalvas.',
    '20 testes automatizados passaram; 25 páginas HTML e recursos locais percorridos via HTTP.',
    'Alinhar a narrativa da IdeiaAI ao comportamento atual e ensaiar com contas/dados preparados.'
]], [2.05,2.35,2.65])
para('Conclusão executiva', 'Heading 1')
para('Não foi encontrada falha crítica reproduzível nesta rodada. Os fluxos de pedidos/estoque e edição de dados empresariais têm testes automatizados úteis. A principal exposição na apresentação é a discrepância entre a promessa de análise por IA e a implementação atual, que devolve conteúdo e valores fixos por categoria. A cobertura automatizada não representa todos os fluxos do produto, e não houve ensaio visual em navegadores/tamanhos variados nem validação de produção.')
para('Legenda de gravidade', 'Heading 2')
para('Alta: pode comprometer credibilidade ou resultado central da demonstração. Média: risco relevante de falha/regressão ou dependência que requer preparação. Baixa: melhoria de qualidade com impacto limitado no roteiro de amanhã.')

para('1. Escopo e método', 'Heading 1')
for x in [
    'Inspeção das rotas Flask, páginas HTML, scripts JavaScript, estilos CSS, configuração e estrutura SQLite.',
    'Execução da suíte existente de BackEnd/tests: unittest discover -s BackEnd/tests.',
    'Smoke test HTTP das 25 páginas; páginas públicas retornaram 200 e páginas protegidas redirecionaram usuários anônimos (302). Recursos locais em Style, Script e Assets retornaram 200.',
    'Revisão estática de autenticação/sessão, CSRF, recuperação de senha, carrinho/pedidos, orçamento e serviços externos citados no código.',
    'Leitura da documentação de execução local e requisitos declarados para configuração de produção.'
]: bullet(x)
para('Como interpretar o resultado', 'Heading 2')
para('O smoke test confirma resposta HTTP e controle básico de acesso anônimo; não equivale à navegação completa de cada formulário. A análise estática identifica comportamento provável pelo código, mas não substitui teste de uso. Os achados abaixo diferenciam evidência observada, inferência de risco e ação recomendada.')

para('2. Fluxos do sistema e cobertura observada', 'Heading 1')
para('As rotas e páginas encontradas mostram dois percursos principais, cliente e empresa. “Não exercitado de ponta a ponta” significa que o fluxo foi mapeado no código, mas não validado integralmente por automação nesta rodada.')
table(['ÁREA / FLUXO','O QUE FOI IDENTIFICADO','COBERTURA NESTA AUDITORIA'],[
('Acesso e conta','Cadastro cliente/empresa, login, sessão, logout, recuperação e redefinição de senha.','Páginas públicas no smoke; estrutura/CSRF inspecionados. Sem suíte ampla de cadastro/login/recuperação.'),
('Perfil e endereço','Visualização/edição de perfil, endereço e consulta de frete.','Rotas identificadas e controle de acesso inspecionado; percurso completo não exercitado.'),
('Vitrine e compra','Catálogo, carrinho, consulta/uso de endereço, frete, criação e acompanhamento de pedidos.','Testes cobrem criação, estoque, agrupamento por empresa, status e cancelamento; fluxo de tela completo não exercitado.'),
('Assistência e coleta','Solicitação de atendimento, descrição do problema, modalidade loja/coleta e consulta de CEP.','Rotas/scripts inspecionados; dependência ViaCEP identificada; sem teste ponta a ponta.'),
('Análise e orçamento','Análise preliminar, geração/apresentação de faixa e prazo, envio da solicitação e decisão do cliente.','Endpoint estático examinado; caminho completo cliente-empresa-cliente sem testes automatizados identificados.'),
('Acompanhamento técnico','Acompanhamento de ordens, preferências de entrega e mudança de etapas.','Rotas e páginas identificadas; sem suíte automatizada ampla de transições nesta rodada.'),
('Operação empresarial','Painel, dados da empresa, aparelhos, solicitações, orçamentos, produtos, pedidos e ordens de serviço.','Testes específicos de dados empresariais e gestão de pedidos; demais páginas privadas redirecionaram anonimamente.'),
('Interface e responsividade','HTML, CSS, recursos e composição de páginas.','Recursos locais retornaram 200. Não foi feito ensaio visual em viewport/navegadores variados nem auditoria de acessibilidade.')
], [1.15,2.45,3.45])

para('3. Achados e gravidade', 'Heading 1')
para('Alta | IdeiaAI é apresentada como análise por IA, mas a resposta é fixa', 'Heading 2')
para('Evidência: POST /api/orcamento/analisar em BackEnd/app.py escolhe mensagens, faixas de preço e prazos a partir de dicionários indexados pela categoria. A rota não consulta BackEnd/ia_prompt.py nem um provedor de modelo. A tela/script descreve o processo como análise IdeiaAI. As faixas são atribuídas genericamente a referências brasileiras de 2026, sem referências verificáveis apresentadas ao usuário.')
para('Risco: um público técnico pode perguntar qual modelo é usado, como se controla a resposta e quais fontes sustentam os valores. O comportamento do produto não sustenta a afirmação de inferência por IA, e a precisão aparente dos valores pode ser interpretada como dado validado.')
para('Ação para amanhã: apresentar como “triagem preliminar” ou “estimativa orientativa por categoria”, explicando que preço e prazo são referências ilustrativas a confirmar pela assistência. Se a narrativa exigir IA real, a integração e validação precisam ocorrer antes de manter essa afirmação.')

para('Média | Cobertura automatizada é estreita para a quantidade de percursos', 'Heading 2')
para('Evidência: os 20 testes concentram-se em pedidos/estoque e dados empresariais. Não foi localizada cobertura automatizada equivalente para cadastro, login, recuperação, perfil/endereço, frete, atendimento/coleta, orçamento e decisão do cliente, ordens e transições de status, responsividade ou acessibilidade.')
para('Risco: áreas sem testes podem regredir sem que a suíte atual indique problema. Ação: para a demo, percorrer manualmente o roteiro escolhido e conferir as contas; depois, ampliar testes de integração nos fluxos de orçamento, atendimento e autenticação.')

para('Média | Serviços externos podem degradar a experiência da demo', 'Heading 2')
para('Evidência: a busca de CEP de coleta chama o ViaCEP no navegador e algumas páginas carregam Google Fonts. A inspeção confirmou apenas os recursos locais, não a disponibilidade dessas dependências externas.')
para('Risco e ação: rede instável pode impedir a consulta do CEP ou mudar a tipografia. Ensaiar no local/rede do evento, testar preenchimento manual e deixar dados de demonstração prontos.')

para('Baixa | Validação de e-mail é permissiva em partes do cadastro', 'Heading 2')
para('Evidência: validações de servidor usam presença de “@” como verificação básica em alguns caminhos. Risco: endereços malformados podem persistir e falhar em recuperação ou comunicação. Melhorar com validação de formato consistente e confirmação por e-mail quando o serviço estiver ativo.')

para('Baixa | Ambiente de produção não comprovado', 'Heading 2')
para('Evidência: COMO_RODAR.md descreve a execução local como desenvolvimento e lista requisitos de produção como HTTPS, chave persistente, SMTP e APP_ENV=production. Não foram fornecidos ambiente de deploy ou ensaio de produção. Isso não bloqueia uma demo local preparada; não permite afirmar prontidão para operação pública.')

para('4. Segurança, dados e estrutura', 'Heading 1')
for x in [
    'Observado no código: token CSRF obrigatório em métodos de escrita, comparação segura, cookies de sessão HttpOnly e SameSite=Lax, duração de sessão definida, proteção de áreas empresariais e controle de acesso associado à conta.',
    'A configuração define cookie Secure quando APP_ENV=production e impede recuperação local nesse modo; a operação em produção ainda depende de variáveis e HTTPS conforme a documentação.',
    'O schema SQLite organiza usuários, empresas, produtos, pedidos/itens, solicitações, orçamentos, ordens e histórico, com chaves estrangeiras e restrições de status/valores.',
    'Testes de pedidos verificam controle de acesso/CSRF, rejeição de carrinho inválido, baixa e reversão de estoque, sequência de status e atomicidade em falha SQL.',
    'Testes de dados empresariais verificam acesso restrito, isolamento por empresa autenticada, CSRF, dados inválidos e duplicidade de CNPJ.'
]: bullet(x)
para('Limite: esta é uma revisão direcionada, não um pentest. Não foi executado teste de carga, revisão completa de acessibilidade, validação SMTP nem avaliação de implantação.')

para('5. Execução dos testes e inspeção de páginas', 'Heading 1')
table(['VERIFICAÇÃO','RESULTADO','OBSERVAÇÃO'],[
('Suíte automatizada existente','20 aprovados','unittest discover -s BackEnd/tests'),
('Primeira tentativa no sandbox','Falha ao criar banco temporário','O %TEMP% do sandbox não permitiu escrita; não foi erro funcional dos testes.'),
('Reexecução com TEMP/TMP em .test-tmp','20 aprovados','Mesma suíte concluída com sucesso.'),
('Páginas HTML','25 percorridas','Páginas públicas 200; páginas protegidas retornaram 302 sem sessão.'),
('CSS, JS e Assets locais','Sem erros HTTP encontrados','Todos os recursos locais percorridos retornaram 200.'),
('UI em múltiplos viewports','Não executada','Exige inspeção visual real em navegadores/tamanhos variados.')
], [1.7,1.7,3.65])

para('6. Melhorias pontuais priorizadas para o MVP', 'Heading 1')
table(['PRIORIDADE','AÇÃO','VALOR PARA A APRESENTAÇÃO'],[
('Antes de apresentar','Ajustar a fala/rótulo da IdeiaAI e deixar explícito que preço/prazo são indicativos.','Evita prometer uma integração ou precisão que o sistema não oferece.'),
('Antes de apresentar','Ensaiar o percurso com conta cliente e conta empresa preparadas; validar solicitação, orçamento, decisão e acompanhamento.','Reduz risco de bloqueio por cadastro, sessão ou ausência de dados.'),
('Antes de apresentar','Ensaiar na rede/local do evento; testar CEP manual e tipografia carregada.','Diminui dependência de serviço externo durante a demonstração.'),
('Se houver tempo','Conferir estoque, produtos ativos, imagens e pedidos de demonstração; usar dados fictícios.','Ajuda a evitar telas vazias e dados pessoais expostos.'),
('Após a apresentação','Adicionar testes de integração para atendimento, orçamento, recuperação, endereço/frete e ordens de serviço.','Dá confiança aos fluxos ainda descobertos pela suíte atual.'),
('Antes de uso público','Validar implantação com HTTPS, chave persistente, SMTP, recuperação e estratégia de backup.','Fecha lacunas entre demo local e serviço operacional.')
], [1.25,3.1,2.7])

para('7. Roteiro sugerido para a demonstração', 'Heading 1')
for x in [
    'Entrar com conta empresarial preparada e mostrar painel e indicadores.',
    'Como cliente, abrir solicitação de assistência, preencher problema e obter a triagem preliminar.',
    'Alternar para a empresa, apresentar a solicitação e registrar/propor o orçamento disponível no sistema.',
    'Voltar ao cliente para demonstrar decisão e acompanhamento, usando registros previamente preparados.',
    'Mostrar vitrine, carrinho e pedido se o estoque e as contas estiverem conferidos.',
    'Manter servidor local iniciado, navegador na página inicial e plano simples de retorno caso a rede externa falhe.'
]: bullet(x)

para('8. Limites e parecer final', 'Heading 1')
para('A auditoria não exercitou cada formulário/transição pela interface, nem realizou teste visual multiplataforma, carga, pentest, auditoria completa de acessibilidade, teste SMTP ou deploy. O resultado é uma avaliação de prontidão baseada nos testes existentes, smoke test HTTP e inspeção de código/configuração. Portanto, não se deve interpretar os 20 testes aprovados como cobertura integral.')
para('Parecer: o MVP pode ser demonstrado localmente com roteiro guiado e dados previamente preparados. Corrigir a narrativa da IdeiaAI e ensaiar o fluxo que será mostrado são as medidas de maior impacto para amanhã à tarde.')

footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('IdeiaTech  |  Auditoria de prontidão do MVP  |  22 out 2026'); r.font.name='Aptos'; r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(gray)
doc.core_properties.title='Auditoria do sistema IdeiaTech para apresentação do MVP'
doc.core_properties.subject='Revisão funcional, técnica, estrutural e de interface'
doc.core_properties.author='IdeiaTech'
doc.save(OUT)
print(OUT)
